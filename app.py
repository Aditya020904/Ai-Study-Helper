import streamlit as st
import tempfile
import requests

# PDF
from langchain_community.document_loaders import PyPDFLoader

# PPT
from pptx import Presentation

# Embeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings

# ================= OLLAMA FUNCTION =================
def generate_response(prompt):
    try:
        res = requests.post(
            "http://localhost:11434/api/generate",
            json={
                "model": "llama3",
                "prompt": prompt,
                "stream": False
            }
        )
        return res.json()["response"]

    except Exception as e:
        return f"❌ Ollama Error: {str(e)}"


# ================= PPT TEXT EXTRACTION =================
def extract_text_from_ppt(ppt_path):

    prs = Presentation(ppt_path)

    text = ""

    for slide in prs.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text


# ================= UI =================
st.set_page_config(
    page_title="AI Study Helper",
    layout="wide"
)

st.title("📚 AI Study Helper (Offline AI)")

# ================= SESSION =================
if "db" not in st.session_state:
    st.session_state.db = None

if "docs" not in st.session_state:
    st.session_state.docs = None

if "history" not in st.session_state:
    st.session_state.history = []

if "quiz" not in st.session_state:
    st.session_state.quiz = ""


# ================= SIDEBAR =================
st.sidebar.title("📂 Upload Files")

uploaded_file = st.sidebar.file_uploader(
    "Upload PDF or PPT",
    type=["pdf", "pptx"]
)

# ================= PROCESS FILE =================
if uploaded_file and st.session_state.db is None:

    file_ext = uploaded_file.name.split(".")[-1].lower()

    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=f".{file_ext}"
    ) as tmp:

        tmp.write(uploaded_file.read())

        file_path = tmp.name

    with st.spinner("Processing File..."):

        # ================= PDF =================
        if file_ext == "pdf":

            loader = PyPDFLoader(file_path)

            docs = loader.load()

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=400,
                chunk_overlap=50
            )

            chunks = splitter.split_documents(docs)

        # ================= PPT =================
        elif file_ext == "pptx":

            ppt_text = extract_text_from_ppt(file_path)

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=400,
                chunk_overlap=50
            )

            chunks = splitter.create_documents([ppt_text])

        # ================= EMBEDDINGS =================
        embeddings = HuggingFaceEmbeddings(
            model_name="sentence-transformers/all-MiniLM-L6-v2"
        )

        db = FAISS.from_documents(
            chunks,
            embeddings
        )

        st.session_state.db = db
        st.session_state.docs = chunks

    st.success("✅ File Ready!")

# ================= SUMMARY =================
st.subheader("📄 File Summary")

if st.button("🚀 Generate Summary"):

    if not st.session_state.docs:

        st.error("Upload file first")

    else:

        context = "\n".join(
            [d.page_content for d in st.session_state.docs[:5]]
        )

        prompt = f"""
        Summarize this document:

        - Key Topics
        - Important Concepts
        - Short Explanation

        {context}
        """

        result = generate_response(prompt)

        st.write(result)

# ================= CHAT =================
st.subheader("💬 Ask Questions")

query = st.text_input(
    "Ask from uploaded file or general"
)

if query:

    context = ""

    if st.session_state.db:

        retriever = st.session_state.db.as_retriever()

        docs = retriever.get_relevant_documents(query)

        context = "\n".join(
            [d.page_content for d in docs[:3]]
        )

    # ================= PROMPT =================
    if context.strip() == "":

        prompt = f"""
        Answer clearly:

        {query}
        """

    else:

        prompt = f"""
        Answer using context.

        Context:
        {context}

        Question:
        {query}
        """

    result = generate_response(prompt)

    st.session_state.history.append(
        (query, result)
    )

    st.write("### 📌 Answer")

    st.write(result)

# ================= QUIZ =================
st.subheader("📝 Quiz Generator")

# Number of Questions
num = st.slider(
    "Number of Questions",
    3,
    15,
    5
)

# Topic Input
quiz_topic = st.text_input(
    "Enter Topic for Quiz",
    placeholder="Example: NLP, Python, DBMS"
)

# Difficulty
difficulty = st.selectbox(
    "Difficulty",
    ["Easy", "Medium", "Hard"]
)

# Generate Quiz
if st.button("Generate Quiz"):

    if not st.session_state.docs and quiz_topic.strip() == "":

        st.error("Upload file or enter topic")

    else:

        context = ""

        # From Uploaded File
        if st.session_state.docs:

            context = "\n".join(
                [d.page_content for d in st.session_state.docs[:5]]
            )

        # From Topic
        if quiz_topic.strip() != "":

            context += f"\nGenerate quiz on topic: {quiz_topic}"

        prompt = f"""
        Create {num} MCQs.

        Difficulty Level: {difficulty}

        Format:

        Q1.
        a)
        b)
        c)
        d)

        Answer:

        Content:
        {context}
        """

        st.session_state.quiz = generate_response(prompt)

# ================= SHOW QUIZ =================
if st.session_state.quiz:

    st.write("### 🧠 Generated Quiz")

    st.write(st.session_state.quiz)

# ================= TEXT SUMMARY =================
st.subheader("🧾 Syllabus / Text Summary")

text = st.text_area(
    "Paste syllabus or notes here"
)

if st.button("Summarize Text"):

    if text.strip() == "":

        st.warning("Enter text")

    else:

        result = generate_response(
            f"Summarize this:\n{text}"
        )

        st.write(result)

# ================= DASHBOARD =================
st.sidebar.markdown("---")

st.sidebar.subheader("📊 Dashboard")

st.sidebar.write(
    f"Questions Asked: {len(st.session_state.history)}"
)

# ================= HISTORY =================
st.sidebar.markdown("---")

st.sidebar.subheader("📜 History")

for q, a in st.session_state.history[::-1]:

    st.sidebar.write(f"Q: {q}")

    st.sidebar.write(f"A: {a[:80]}...")

    st.sidebar.write("---")

# ================= CLEAR =================
if st.sidebar.button("🗑 Clear History"):

    st.session_state.history = []
    st.session_state.quiz = ""
    st.session_state.db = None
    st.session_state.docs = None

    st.success("✅ Cleared!")